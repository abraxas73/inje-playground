# -*- coding: utf-8 -*-
"""제안도우미-일반현황 엔진 — 합성 일반현황·고객사 템플릿으로 두 방식을 검증한다(실제 회사 자료는 픽스처에 넣지 않는다)."""
import struct
import zipfile
import zlib
from collections import Counter

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

from service import proposal as P


def _png(w=40, h=20):
    raw = b"".join(b"\x00" + b"\x30\x60\xc0" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def make_profile(tmp_path):
    """16:9 일반현황 3장: ① 제목+브랜드 파랑 글 ② 표 ③ 그림+차트."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    lay = prs.slide_layouts[5]  # Title Only
    s1 = prs.slides.add_slide(lay)
    s1.shapes.title.text = "회사 개요"
    box = s1.shapes.add_textbox(Cm(2), Cm(5), Cm(20), Cm(3))
    box.text_frame.text = "설립 2007년 · 임직원 312명"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = RGBColor(0x00, 0x6E, 0xFF)
    run.font.size = Pt(20)
    run.font.name = "Pretendard"
    s2 = prs.slides.add_slide(lay)
    s2.shapes.title.text = "주요 사업 실적"
    t = s2.shapes.add_table(3, 3, Cm(2), Cm(5), Cm(28), Cm(4)).table
    for r, row in enumerate([["연도", "고객", "금액"], ["2025", "A공사", "12억"], ["2026", "B은행", "8억"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    s3 = prs.slides.add_slide(lay)
    s3.shapes.title.text = "재무 현황"
    img = tmp_path / "a.png"
    img.write_bytes(_png())
    s3.shapes.add_picture(str(img), Cm(2), Cm(5), Cm(8), Cm(4))
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("매출", (180, 210))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(12), Cm(5), Cm(18), Cm(10), data)
    path = tmp_path / "profile.pptx"
    prs.save(str(path))
    return path


@pytest.fixture
def profile(tmp_path):
    return make_profile(tmp_path)


def _template(tmp_path, wide, name):
    prs = Presentation()
    if wide:
        prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "고객사 예시 장표"
    path = tmp_path / name
    prs.save(str(path))
    return path


def _ids_unique(slide):
    ids = [e.get("id") for e in slide.shapes._spTree.iter(qn("p:cNvPr"))]
    return not [i for i, n in Counter(ids).items() if n > 1]


@pytest.mark.parametrize("wide", [True, False])
def test_restyle_keeps_content_and_takes_customer_look(profile, tmp_path, wide):
    tpl = _template(tmp_path, wide, f"tpl-{wide}.pptx")
    out = tmp_path / "out.pptx"
    res = P.restyle(profile, tpl, out)
    prs = Presentation(str(out))
    t = Presentation(str(tpl))
    assert res["slides"] == 3 and len(prs.slides) == 3, "예시 장표는 빠지고 일반현황 장 수만큼"
    assert (prs.slide_width, prs.slide_height) == (t.slide_width, t.slide_height)
    assert [s.shapes.title.text for s in prs.slides] == ["회사 개요", "주요 사업 실적", "재무 현황"]
    xml = "".join(etree_str(s) for s in prs.slides)
    assert "설립 2007년" in xml and "B은행" in xml and "8억" in xml, "숫자·표 내용은 원본 그대로"
    assert 'val="006EFF"' not in xml and 'schemeClr val="accent1"' in xml, "브랜드 파랑 → 고객사 accent1"
    assert 'typeface="+mn-lt"' in xml and "Pretendard" not in xml, "글꼴은 고객사 테마 글꼴"
    s3 = prs.slides[2]
    assert any(sh.shape_type is not None for sh in s3.shapes)
    rels = [r.reltype for r in s3.part.rels.values()]
    assert any(r.endswith("/image") for r in rels) and any(r.endswith("/chart") for r in rels), "그림·차트 관계도 옮긴다"
    for s in prs.slides:
        assert _ids_unique(s), "도형 id 중복은 PowerPoint 복구 경고"
        for sh in s.shapes:  # 본문 영역(슬라이드 안)에 들어간다
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= prs.slide_width + 1 and sh.top + sh.height <= prs.slide_height + 1
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert len(names) == len(set(names)), "partname 충돌 없음"


def etree_str(slide):
    from lxml import etree
    return etree.tostring(slide.shapes._spTree, encoding="unicode")


def test_reflow_fills_customer_layouts(profile, tmp_path):
    tpl = _template(tmp_path, True, "tpl.pptx")
    info = P.template_info(P.open_pptx(tpl))
    body_layout = next(l["index"] for l in info["layouts"] if [p["kind"] for p in l["placeholders"]].count("body") == 1 and "title" in [p["kind"] for p in l["placeholders"]])
    outline = P.profile_outline(P.open_pptx(profile))
    assert outline[1]["tables"][0][2] == ["2026", "B은행", "8억"]
    assert outline[2]["images"][0]["id"] == "s3-1"
    spec = {"slides": [
        {"layout": body_layout, "title": "회사 개요", "body": ["설립 2007년", {"text": "임직원 312명", "level": 1}]},
        {"layout": body_layout, "title": "주요 사업 실적", "table": {"header": ["연도", "고객", "금액"], "rows": [["2025", "A공사", "12억"], ["2026", "B은행", "8억"]]}},
        {"layout": 99, "title": "재무 현황", "images": ["s3-1", "없는그림"]},
    ]}
    out = tmp_path / "reflow.pptx"
    assert P.reflow(profile, tpl, spec, out)["slides"] == 3
    prs = Presentation(str(out))
    s1, s2, s3 = prs.slides
    assert "임직원 312명" in etree_str(s1)
    tables = [sh for sh in s2.shapes if sh.has_table]
    assert tables and tables[0].table.cell(2, 2).text == "8억"
    assert any(sh.shape_type == 13 for sh in s3.shapes), "그림 1장(없는 id는 무시)"
    for s in prs.slides:  # 채우지 않은 자리표시자는 남기지 않는다
        assert all(ph.text_frame.text.strip() or ph.placeholder_format.type in (1, 3) for ph in s.placeholders if ph.has_text_frame)


def test_potx_and_macro_files(profile, tmp_path):
    tpl = _template(tmp_path, True, "tpl.pptx")
    potx = tmp_path / "tpl.potx"
    with zipfile.ZipFile(tpl) as src, zipfile.ZipFile(potx, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            dst.writestr(item, data)
    assert P.restyle(profile, potx, tmp_path / "o.pptx")["slides"] == 3
    bad = tmp_path / "m.pptm"
    with zipfile.ZipFile(tpl) as src, zipfile.ZipFile(bad, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("ppt/vbaProject.bin", b"x")
    with pytest.raises(P.ProposalError, match="매크로"):
        P.open_pptx(bad)
    (tmp_path / "x.pptx").write_bytes(b"not a zip")
    with pytest.raises(P.ProposalError):
        P.open_pptx(tmp_path / "x.pptx")


def test_endpoints_require_token_and_allowed_hosts(client, auth):
    assert client.post("/proposal/outline", json={"profileUrl": "https://example.supabase.co/a"}).status_code == 401
    assert client.post("/proposal/outline", headers=auth, json={"profileUrl": "https://evil.test/a"}).status_code == 400
    assert client.post("/proposal/outline", headers=auth, json={}).status_code == 400
    body = {"mode": "reflow", "profileUrl": "https://example.supabase.co/a", "templateUrl": "https://example.supabase.co/b", "pptxUrl": "https://example.supabase.co/c"}
    assert client.post("/proposal/build", headers=auth, json=body).status_code == 400, "reflow는 spec 필수"
    assert client.post("/proposal/build", headers=auth, json={**body, "mode": "restyle", "pptxUrl": "https://evil.test/c"}).status_code == 400

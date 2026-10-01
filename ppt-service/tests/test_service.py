def test_health_requires_token(client):
    assert client.get("/health").status_code == 401
    assert client.get("/health", headers={"x-ppt-token": "wrong"}).status_code == 401


def test_health_non_ascii_token_is_401(client):
    assert client.get("/health", headers={"x-ppt-token": "é".encode()}).status_code == 401


def test_health(client, auth):
    r = client.get("/health", headers=auth)
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True and body["templateSlides"] == 106 and body["layouts"] == 69


def test_catalog_shape(client, auth):
    cat = client.get("/catalog", headers=auth).json()
    names = {e["name"] for e in cat["layouts"]}
    assert len(cat["layouts"]) == 69 and "card-4" in names and "free-title" in names
    card4 = next(e for e in cat["layouts"] if e["name"] == "card-4")
    assert card4["arity"] == 4 and card4["closing"] == {"required": True, "maxLines": 2}
    assert card4["example"]["layout"] == "card-4" and len(card4["example"]["cards"]) == 4
    assert card4["example"]["cards"][0]["title"] == "…" and card4["example"]["closing"] == "…"
    assert card4["capacity"]["title"] == [9, 2] and card4["capacity"]["body"] == [17, 4] and "page_title" not in card4["capacity"]
    assert cat["capacityCommon"] == {"page_title": [47, 2], "section_label": [20, 1]}
    assert cat["message"]["capacity"] == {}  # message는 패키지가 용량 검사 대신 권고만 낸다
    assert card4["table"] is None
    note = next(e for e in cat["layouts"] if e["name"] == "table-note")
    assert note["table"] == {"widthCm": 27.0, "heightCm": 7.1, "rowsOneLine": 6, "rowsTwoLine": 4, "charsPerLine": 95}
    full = next(e for e in cat["layouts"] if e["name"] == "table-full")
    assert full["table"]["rowsOneLine"] == 9
    # 마무리 바(bar16)는 양옆 6cm 여백 → 한 줄 26자(패키지 표 47자가 아니라). msg형은 46자
    assert card4["capacity"]["closing"] == [26, 3]
    assert next(e for e in cat["layouts"] if e["name"] == "kpi-3-cards-3")["capacity"]["closing"] == [46, 2]
    image4 = next(e for e in cat["layouts"] if e["name"] == "image-4")
    assert image4["example"]["images"] == ["src:장:번호"] * 4
    free = next(e for e in cat["layouts"] if e["name"] == "free-title")
    assert free["example"] == {"layout": "free-title", "title": ["…", "…"], "source": {"slide": 1}}
    assert cat["message"]["example"] == {"layout": "message", "headline": "…", "detail": "…"}
    assert cat["products"] == ["aicubeit", "devopsit", "openstackit", "secloudit", "tabcloudit"]
    assert cat["overview"] == ["lineup", "tafa", "tafa-layers"]


from pathlib import Path


def make_source_pptx(path: Path):
    from pptx import Presentation
    from pptx.util import Cm, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    tb = s1.shapes.add_textbox(Cm(1), Cm(1), Cm(20), Cm(2))
    tb.text_frame.text = "클라우드 전환 배경"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    body = s1.shapes.add_textbox(Cm(1), Cm(5), Cm(20), Cm(5))
    body.text_frame.text = "첫 문단\n둘째 문단"
    s2 = prs.slides.add_slide(blank)
    from PIL import Image  # Pillow는 python-pptx 의존성으로 함께 설치된다
    img = path.parent / "pic.png"
    Image.new("RGB", (40, 30), "blue").save(img)
    s2.shapes.add_picture(str(img), Cm(2), Cm(2), Cm(10), Cm(6))
    prs.save(str(path))


def test_extract_slides(tmp_path):
    from service.extract import extract_slides
    src = tmp_path / "src.pptx"
    make_source_pptx(src)
    slides = extract_slides(str(src))
    assert [s["no"] for s in slides] == [1, 2]
    assert slides[0]["title"] == "클라우드 전환 배경"
    assert slides[0]["texts"] == ["클라우드 전환 배경", "첫 문단\n둘째 문단"]
    assert slides[0]["pictures"] == 0 and slides[0]["titleBottomCm"] == 3.0
    assert slides[1]["texts"] == [] and slides[1]["pictures"] == 1 and slides[1]["title"] is None


def test_extract_endpoint(client, auth, tmp_path, monkeypatch):
    import service.storage as S
    src = tmp_path / "src.pptx"
    make_source_pptx(src)
    def fake_download(url, dest):
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(src.read_bytes())
        return dest
    monkeypatch.setattr(S, "download", fake_download)
    r = client.post("/extract", headers=auth, json={"sourceUrl": "https://example.supabase.co/storage/v1/object/sign/ppt/source/a.pptx?token=x"})
    assert r.status_code == 200 and len(r.json()["slides"]) == 2


def test_extract_rejects_foreign_host(client, auth):
    r = client.post("/extract", headers=auth, json={"sourceUrl": "https://evil.example.com/a.pptx"})
    assert r.status_code == 400


URL_OK = "https://example.supabase.co/storage/v1/object/sign/ppt/source/a.pptx?token=x"


def test_check_body_rejects_non_integer_length():
    import pytest
    import main
    from fastapi import HTTPException

    class Stub:
        def __init__(self, length):
            self.headers = {"content-length": length}

    with pytest.raises(HTTPException) as e:
        main.check_body(Stub("abc"))
    assert e.value.status_code == 400 and e.value.detail == "invalid content-length"
    with pytest.raises(HTTPException) as e:
        main.check_body(Stub(str(main.MAX_BODY + 1)))
    assert e.value.status_code == 413


def test_extract_download_failure_is_502(client, auth, monkeypatch):
    import httpx
    import service.storage as S

    def boom(url, dest):
        raise httpx.ConnectError("boom")
    monkeypatch.setattr(S, "download", boom)
    r = client.post("/extract", headers=auth, json={"sourceUrl": URL_OK})
    assert r.status_code == 502


def test_extract_corrupt_pptx_is_400(client, auth, monkeypatch):
    import service.storage as S

    def junk(url, dest):
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(b"not a pptx")
        return dest
    monkeypatch.setattr(S, "download", junk)
    r = client.post("/extract", headers=auth, json={"sourceUrl": URL_OK})
    assert r.status_code == 400 and r.json()["detail"] == "pptx 파일을 열 수 없습니다(손상되었거나 pptx가 아닙니다)"


MINI_SPEC = {
    "meta": {"title": ["최소 예제로 확인하는", "빌드 파이프라인입니다."], "subtitle": "테스트", "ver": "01", "date": "2026. 09. 30", "dept": "클라우드네이티브센터"},
    "sections": [{"name": "개요", "slides": [{
        "layout": "card-3",
        "title": ["세 가지 이유로 정리하는", "표준화가 필요한 배경입니다."],
        "cards": [
            {"title": "속도", "body": ["배포가 수작업이라 실수가 반복됨", "환경마다 설정이 갈리는 것"]},
            {"title": "비용", "body": ["장비 교체에 인력이 묶이는 구조"]},
            {"title": "품질", "body": ["장애 원인 파악까지 평균 네 시간"]},
        ],
        "closing": "서버가 아니라 [[일하는 방식]]을 바꾸는 것이 목적입니다.",
    }]}],
}


def test_build_deck(tmp_path):
    from service.builder import build_deck
    res = build_deck(MINI_SPEC, tmp_path, None, None)
    assert res["slides"] == 5            # 표지·목차·간지·본문·뒷표지
    assert res["issues"] == {}
    assert (tmp_path / "deck.pptx").stat().st_size > 1_000_000
    assert "layout: card-3" in (tmp_path / "deck.yaml").read_text(encoding="utf-8")
    assert isinstance(res["advisories"], list)


def test_build_error_has_position(tmp_path):
    from service.builder import BuildError, build_deck
    bad = {**MINI_SPEC, "sections": [{"name": "개요", "slides": [{**MINI_SPEC["sections"][0]["slides"][0], "layout": "card-4"}]}]}
    try:
        build_deck(bad, tmp_path, None, None)
        assert False, "should raise"
    except BuildError as e:
        assert e.kind == "spec" and e.section == 0 and e.slide == 0 and "card-4" in e.message


def test_build_endpoint(client, auth, tmp_path, monkeypatch):
    import service.storage as S
    uploaded = {}
    monkeypatch.setattr(S, "upload", lambda url, path, ct: uploaded.setdefault(url, path.stat().st_size))
    body = {"spec": MINI_SPEC, "upload": {"pptxUrl": "https://example.supabase.co/storage/v1/object/upload/sign/ppt/decks/d/v1/deck.pptx?token=a",
                                          "yamlUrl": "https://example.supabase.co/storage/v1/object/upload/sign/ppt/decks/d/v1/deck.yaml?token=b"}}
    r = client.post("/build", headers=auth, json=body)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["ok"] is True and j["slides"] == 5 and len(uploaded) == 2 and j["bytes"] > 1_000_000


def test_build_endpoint_spec_error(client, auth):
    bad = {**MINI_SPEC, "sections": [{"name": "개요", "slides": [{"layout": "no-such-layout", "title": ["a", "b."]}]}]}
    r = client.post("/build", headers=auth, json={"spec": bad, "upload": {"pptxUrl": "https://example.supabase.co/x?token=1", "yamlUrl": "https://example.supabase.co/y?token=2"}})
    assert r.status_code == 422
    j = r.json()
    assert j["ok"] is False and j["kind"] == "spec" and j["section"] == 0 and j["slide"] == 0 and "no-such-layout" in j["message"]


def _free_spec(slide):
    return {**MINI_SPEC, "sections": [{"name": "개요", "slides": [slide]}]}


FREE = {"layout": "free-title", "title": ["원고를 옮겨 붙인", "자유 배치 장표입니다."]}


def test_build_yaml_has_no_work_paths(tmp_path):
    import copy
    from service.builder import build_deck
    src = tmp_path / "src.pptx"
    make_source_pptx(src)
    spec = _free_spec({**FREE, "images": ["src:2:1"], "source": {"slide": 1}})
    orig = copy.deepcopy(spec)
    res = build_deck(spec, tmp_path / "w", src, None)
    text = (tmp_path / "w" / "deck.yaml").read_text(encoding="utf-8")
    assert "/tmp" not in text and "ppt-" not in text and str(tmp_path) not in text
    assert "src:2:1" in text and spec == orig and res["slides"] == 5


def _post_spec(client, auth, spec):
    return client.post("/build", headers=auth, json={"spec": spec, "upload": {"pptxUrl": "https://example.supabase.co/x?token=1", "yamlUrl": "https://example.supabase.co/y?token=2"}})


def test_build_meta_title_type_error(client, auth):
    r = _post_spec(client, auth, {**MINI_SPEC, "meta": {**MINI_SPEC["meta"], "title": 123}})
    j = r.json()
    assert r.status_code == 422 and j["kind"] == "spec" and j["section"] is None


def test_build_attribute_error_has_position(client, auth):
    r = _post_spec(client, auth, _free_spec({"layout": "product", "product": 123, "title": ["a", "b."]}))
    j = r.json()
    assert r.status_code == 422 and j["kind"] == "spec" and (j["section"], j["slide"]) == (0, 0)


def test_build_bad_source_slide_has_position(client, auth, tmp_path, monkeypatch):
    import service.storage as S
    src = tmp_path / "src.pptx"
    make_source_pptx(src)
    def fake_download(url, dest):
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(src.read_bytes())
        return dest
    monkeypatch.setattr(S, "download", fake_download)
    body = {"spec": _free_spec({**FREE, "source": {"slide": "abc"}}), "sourceUrl": "https://example.supabase.co/s.pptx?token=1",
            "upload": {"pptxUrl": "https://example.supabase.co/x?token=1", "yamlUrl": "https://example.supabase.co/y?token=2"}}
    r = client.post("/build", headers=auth, json=body)
    j = r.json()
    assert r.status_code == 422 and j["kind"] == "spec" and (j["section"], j["slide"]) == (0, 0)


def test_build_storage_connect_error(client, auth, monkeypatch):
    import httpx
    import service.storage as S
    def boom(url, dest):
        raise httpx.ConnectError("boom")
    monkeypatch.setattr(S, "download", boom)
    body = {"spec": MINI_SPEC, "sourceUrl": "https://example.supabase.co/s.pptx?token=1",
            "upload": {"pptxUrl": "https://example.supabase.co/x?token=1", "yamlUrl": "https://example.supabase.co/y?token=2"}}
    r = client.post("/build", headers=auth, json=body)
    assert r.status_code == 502 and r.json()["kind"] == "internal" and r.json()["message"] == "스토리지 연결 오류"


def test_strip_table_accents_removes_markers_only_in_table_cells():
    from service.builder import _strip_table_accents
    spec = {"sections": [{"slides": [{"layout": "table", "closing": "[[남긴다]]", "tables": [{"header": ["구분", "[[LLM Wiki]]"], "rows": [["합성 시점", "[[넣을 때 한 번]]"], ["비용", 3]]}]}]}]}
    _strip_table_accents(spec)
    t = spec["sections"][0]["slides"][0]["tables"][0]
    assert t["header"] == ["구분", "LLM Wiki"] and t["rows"] == [["합성 시점", "넣을 때 한 번"], ["비용", 3]]
    assert spec["sections"][0]["slides"][0]["closing"] == "[[남긴다]]"

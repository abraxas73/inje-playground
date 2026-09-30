#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""템플릿에서 슬롯 용량을 재서 innogrid_ppt/capacity.py 를 다시 만든다.

글자수 상한을 손으로 어림하지 않는다. 상자의 **실제 폭·높이·글자 크기**에서
"한 줄에 몇 자 / 몇 줄" 을 계산해 박아 넣는다. 템플릿이 바뀌면 이 스크립트를
다시 돌리면 된다.

    python tools/measure.py

한글 한 글자를 1em으로, 줄 높이는 상자의 실제 줄간격(없으면 1.3em)으로 본다.
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation                      # noqa: E402

from innogrid_ppt import slots as S                # noqa: E402
from innogrid_ppt import tokens as T               # noqa: E402
from innogrid_ppt import reflow                    # noqa: E402
from innogrid_ppt.geom import walk_abs             # noqa: E402

LINE_RATIO = 1.3          # 줄 높이 / 글자 크기 (템플릿 줄간격 130%)
DEFAULT_PT = 9.0


def font_pt(shape, fallback=DEFAULT_PT):
    for para in shape.text_frame.paragraphs:
        for r in para.runs:
            if r.font.size:
                return r.font.size.pt
    return fallback


def capacity(node, pt=None):
    """이 상자에 들어가는 (줄당 글자 수, 줄 수). 계산식은 slots.capacity_of() 하나다 —
    빌더가 8pt 축소를 판정할 때도 같은 식을 쓴다."""
    return S.capacity_of(node, pt)


def slot_iter(spec):
    """스펙 안의 모든 (역할, 좌표) 를 훑는다."""
    for g in (spec.get("groups") or []) + (spec.get("rows_groups") or []):
        for role, yv in g["slots"].items():
            yield role, yv
    for role, yv in (spec.get("single") or {}).items():
        yield role, yv
    for key in ("main_body",):
        if key in spec:
            yield key, spec[key]
    box = spec.get("box")
    if box:
        yield "note_title", box["title"]
        yield "note_body", box["body"]
    c = spec.get("closing")
    if c:
        yield "closing", c["y"]


def main():
    from innogrid_ppt.template import find_template
    tpl = find_template(root=ROOT)
    prs = Presentation(str(tpl))
    out = {}

    for name, spec in T.LAYOUTS.items():
        if spec.get("free"):
            continue
        slide = prs.slides[spec["slide"] - 1]
        # 단 수를 바꿔 쓰는 장표(card-5·6)는 칸이 좁아진 만큼 줄당 글자 수가 준다.
        # 좌우 여백은 고정이므로 안쪽 폭 비율로 환산한다 (가이드 18·19).
        regrid = spec.get("regrid")
        shrink = ((reflow.CARD_GRID[regrid] - 2 * reflow.CARD_INSET) / 5.50) if regrid else 1.0
        nodes = list(walk_abs(slide.shapes))
        group_roles = {r for g in (spec.get("groups") or []) for r in g["slots"]}
        for role, yv in slot_iter(spec):
            y = yv[0] if isinstance(yv, (tuple, list)) else yv
            xmin = yv[1] if isinstance(yv, (tuple, list)) and len(yv) > 1 else None
            xmax = yv[2] if isinstance(yv, (tuple, list)) and len(yv) > 2 else None
            got = S.band_nodes(slide, y, x_min=xmin, x_max=xmax, _nodes=nodes)
            if not got:
                continue
            # 같은 역할의 상자 중 가장 좁은 것을 기준으로 잡는다 (제일 먼저 넘친다)
            best = None
            for nd in got:
                cap = capacity(nd, font_pt(nd.shape))
                if role in group_roles:                # 카드 안 슬롯만 좁아진다. 전폭 요소는 그대로
                    cap = (max(1, int(cap[0] * shrink)), cap[1])
                if best is None or cap[0] * cap[1] < best[0] * best[1]:
                    best = cap
            key = (name, role)
            if key not in out or best[0] * best[1] < out[key][0] * out[key][1]:
                out[key] = best

        # 페이지 타이틀 · 섹션 라벨
        for role, y in (("page_title", spec.get("title")), ("section_label", spec.get("label"))):
            if y is None:
                continue
            got = S.band_nodes(slide, y, _nodes=nodes)
            if got:
                out[(name, role)] = capacity(got[0], font_pt(got[0].shape, 16.0))

        # 키워드 칩
        chips = spec.get("chips")
        if chips:
            got = S.band_nodes(slide, chips["y"], _nodes=nodes)
            if got:
                out[(name, "chip")] = capacity(got[0], font_pt(got[0].shape, T.CHIP_FONT_PT))

        # 아젠다 행
        ag = spec.get("agenda")
        if ag:
            got = S.band_nodes(slide, ag["ys"][0], _nodes=nodes)
            if len(got) >= 2:
                out[(name, "label")] = capacity(got[0], font_pt(got[0].shape, 10.0))
                out[(name, "text")] = capacity(got[1], font_pt(got[1].shape, 10.0))

        # 좌우 요약 날개 (제목 · 설명)
        wings = spec.get("wings")
        if wings:
            for side in ("left_x", "right_x"):
                if side not in wings:
                    continue
                xmin, xmax = wings[side]
                got = S.band_nodes(slide, wings["ys"][0], x_min=xmin, x_max=xmax, _nodes=nodes)
                if got:
                    out[(name, "wing")] = capacity(got[0], font_pt(got[0].shape, 10.0))
                for by in wings.get("body_ys", [])[:1]:
                    got = S.band_nodes(slide, by, x_min=xmin, x_max=xmax, _nodes=nodes)
                    if got:
                        out[(name, "wing_body")] = capacity(got[0], font_pt(got[0].shape, 9.0))

        # 비교 장표의 요약 칩 상자
        sb = spec.get("sub_boxes")
        if sb:
            got = S.band_nodes(slide, sb["y"], _nodes=nodes)
            if got:
                out[(name, "sub_box")] = capacity(got[0], font_pt(got[0].shape, 9.0))

    lines = [
        "# -*- coding: utf-8 -*-",
        '"""슬롯 용량 — (줄당 글자 수, 줄 수). **tools/measure.py 가 생성한다. 손으로 고치지 말 것.**',
        "",
        "상자의 실제 폭·높이·글자 크기에서 잰 값이다. 한글 한 글자를 1em,",
        f"줄 높이를 {LINE_RATIO}em 으로 본다. 오버플로 판정은 slots.check_fit() 이 한다.",
        '"""',
        "",
        "CAPACITY = {",
    ]
    for (name, role) in sorted(out):
        cpl, ln = out[(name, role)]
        lines.append(f'    ("{name}", "{role}"): ({cpl}, {ln}),')
    lines.append("}")
    path = ROOT / "innogrid_ppt" / "capacity.py"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"{path} — 슬롯 {len(out)}개")


if __name__ == "__main__":
    main()

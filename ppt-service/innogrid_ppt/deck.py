# -*- coding: utf-8 -*-
"""덱 조립.

원본 106장이 든 템플릿을 열고 → 필요한 장표를 뒤에 복제해 쌓고 →
마지막에 원본 106장(디자인 가이드 1~21 포함)과 **템플릿의 슬라이드 구역(section)** 을 잘라낸다.
(프롬프트 참고 사항: "디자인 적용 완료된 PPT에서는 템플릿상의 슬라이드 구역을 제거하고 파일 생성")
"""

import sys

from pptx import Presentation

from . import builders as B
from . import tokens as T
from .clone import clone_slide, drop_slides, drop_sections

TEMPLATE_SLIDES = T.TEMPLATE_SLIDES


def _advise(msg):
    print(f"[권고] {msg}", file=sys.stderr)


class Deck:
    def __init__(self, template_path):
        self.prs = Presentation(template_path)
        B.TEMPLATE_PATH = template_path
        B.TEMPLATE_PRS = self.prs
        if len(self.prs.slides) != TEMPLATE_SLIDES:
            raise ValueError(
                f"템플릿이 {TEMPLATE_SLIDES}장이어야 하는데 {len(self.prs.slides)}장이다.\n"
                f"  INNOGRID_PPT_Template_v1_0_latest.pptx (2026-09 최신본 · 디자인가이드 병합본)가 맞는지 확인할 것.\n"
                f"  86장이면 v1.1, 33장이면 v1.0, 26·22장이면 그 이전 배포본이다 — 번호·좌표가 전부 다르다."
            )
        self._origin = list(self.prs.slides)

    def use(self, n):
        """1-base 원본 번호로 복제."""
        return clone_slide(self.prs, self._origin[n - 1])

    # ── 골격 ──────────────────────────────────────────────────
    def cover(self, **kw):
        B.build_cover(self.use(T.COVER), **kw)

    def toc(self, sections):
        """섹션 최대 8개. 5개를 넘으면 **같은 장의 위쪽에 행을 추가한다** (템플릿 노트 23)."""
        if len(sections) > T.TOC_MAX:
            raise ValueError(f"섹션은 최대 {T.TOC_MAX}개다 ({len(sections)}개가 왔다) — 템플릿 노트")
        B.build_toc(self.use(T.TOC), sections)
        return 1

    def divider(self, index, name, subs=None):
        """하위 섹션이 있으면 25번, 없으면 24번."""
        if subs:
            B.build_divider_sub(self.use(T.DIVIDER_SUB), index=index, name=name, subs=subs)
        else:
            B.build_divider(self.use(T.DIVIDER), index=index, name=name)

    def message(self, **kw):
        B.build_message(self.use(T.MESSAGE), **kw)

    def back_cover(self):
        self.use(T.BACK)

    # ── 본문 ──────────────────────────────────────────────────
    def body(self, name, data):
        B.build_body(self.use(T.LAYOUTS[name]["slide"]), name, data)

    def product(self, product, kind, data):
        p = T.PRODUCTS[product]
        B.build_product(self.use(p[kind]), kind, data)

    def product_overview(self, which):
        self.use(T.PRODUCT_OVERVIEW[which])

    # ── 마무리 ────────────────────────────────────────────────
    def save(self, path):
        drop_slides(self.prs, TEMPLATE_SLIDES)   # 원본 106장 제거 (가이드 1~21 포함)
        drop_sections(self.prs)                  # 템플릿의 슬라이드 구역 제거 (프롬프트 참고 사항)
        self.prs.save(path)
        return path


# ══════════════════════════════════════════════════════════════════
#  장표 지정 풀기
# ══════════════════════════════════════════════════════════════════

_COUNT_KEYS = ("cards", "items", "steps", "panels", "notes", "flow")


def _count(sl):
    for key in _COUNT_KEYS:
        if key in sl and isinstance(sl[key], (list, tuple)):
            return len(sl[key])
    return None


def resolve(sl):
    """장표 지정을 이름으로 푼다. 템플릿 슬라이드 번호(27~92)도 받는다."""
    raw = str(sl.get("layout", "")).strip().lower()
    if not raw:
        raise ValueError("장표에 'layout'이 없다")

    if raw.isdigit():
        n = int(raw)
        if n in T.BY_SLIDE:
            return T.BY_SLIDE[n]
        raise ValueError(
            f"슬라이드 {n}번은 본문 장표가 아니다. 본문은 27~92번이다.\n"
            f"  표지·목차·간지·뒷표지는 meta·sections에서 자동 생성된다")

    if raw in T.ALIASES and T.ALIASES[raw] in T.LAYOUTS:
        raw = T.ALIASES[raw]

    if raw in T.LAYOUTS:
        want = T.LAYOUTS[raw].get("arity")
        got = _count(sl)
        if want is not None and got is not None and got != want:
            raise ValueError(
                f"'{raw}'는 항목 {want}개 고정인데 {got}개가 왔다.\n"
                f"  → {_sibling(raw, got)}")
        return raw

    if raw in T.RENAMED:
        raise ValueError(f"'{raw}'는 이전 템플릿 이름이다.\n  → 최신본에서는 '{T.RENAMED[raw]}'")

    if raw in T.ARITY_FAMILIES:
        raise ValueError(
            f"'{raw}'만으로는 장표가 정해지지 않는다. 단 수는 이름의 일부다.\n"
            f"  → {' / '.join(T.ARITY_FAMILIES[raw])}")

    raise ValueError(
        f"모르는 장표 '{sl['layout']}'.\n"
        f"  쓸 수 있는 이름: {', '.join(sorted(T.LAYOUTS))}")


def _sibling(name, got):
    for names in T.ARITY_FAMILIES.values():
        if name in names:
            same = [o for o in names if T.LAYOUTS[o].get("arity") == got]
            if same:
                return f"'{same[0]}'를 쓸 것 (항목 {got}개)"
    return "항목을 묶거나 장을 분리할 것 (CLAUDE.md §5)"


# ══════════════════════════════════════════════════════════════════
#  덱 사양(dict) → pptx
# ══════════════════════════════════════════════════════════════════

def build(template_path, spec):
    d = Deck(template_path)
    meta = spec.get("meta") or {}
    # 기존 제안서에 끼워 넣을 본문만 뽑을 때 — meta.body_only: true
    body_only = bool(meta.get("body_only"))

    if not body_only:
        if "title" not in meta:
            raise ValueError("meta.title이 없다 (본문만 뽑으려면 meta.body_only: true)")
        # 표지 노트: "부서명은 원고에서 파악되면 넣고, 힘들면 '부서명'이라고만 표기"
        dept = meta.get("dept") or T.DEFAULT_DEPT
        if not meta.get("dept"):
            _advise(f"meta.dept가 없어 표지에 '{T.DEFAULT_DEPT}'로 표기했다 (템플릿 노트 22)")
        d.cover(
            title=meta["title"], subtitle=meta.get("subtitle", ""),
            ver=meta.get("ver", "01"), date=meta.get("date", ""),
            dept=dept, author=meta.get("author", ""),
        )

    sections = spec["sections"]
    if not body_only:
        d.toc([s["name"] for s in sections])

    for i, sec in enumerate(sections):
        subs = sec.get("subs") or []
        if not body_only:
            d.divider(i, sec["name"], subs)
        base = sec.get("label") or f"{i + 1:02d}. {sec['name']}"
        for sl in sec.get("slides", []):
            # 프롬프트: "목차에 하위 목차가 있을 경우 내지 상단의 '01. 섹션명 : 하위섹션명' 구조 준수"
            label = base
            sub = sl.get("sub")
            if sub:
                label = f"{base} : {sub}"
            elif subs and not sec.get("label") and str(sl.get("layout", "")).lower() != "message":
                _advise(f"[{sec['name']}] 하위 섹션이 있는데 장표에 'sub'가 없다 — "
                        f"라벨은 '01. 섹션명 : 하위섹션명' 구조여야 한다 (프롬프트 필수 준수 사항)")
            _emit(d, label, sl)

    if not body_only:
        d.back_cover()
    _lint(spec)
    return d


def _emit(d, label, sl):
    raw = str(sl.get("layout", "")).strip().lower()

    # 핵심 메시지 — 섹션 라벨·타이틀이 없는 독립 장표
    if raw == "message":
        d.message(headline=sl["headline"], detail=sl.get("detail"))
        return

    # 제품 소개 — product: openstackit / kind: intro|features
    if raw.startswith("product"):
        which = sl.get("product", "").strip().lower()
        if which in T.PRODUCT_OVERVIEW:
            d.product_overview(which)
            return
        if which not in T.PRODUCTS:
            raise ValueError(
                f"모르는 제품 '{sl.get('product')}'.\n"
                f"  → {', '.join(sorted(T.PRODUCTS))} 또는 {', '.join(T.PRODUCT_OVERVIEW)}")
        kind = "features" if raw.endswith("features") else "intro"
        data = dict(sl)
        data.setdefault("label", label)
        d.product(which, kind, data)
        return

    name = resolve(sl)
    data = dict(sl)
    data.setdefault("section_label", label)
    d.body(name, data)


def _lint(spec):
    """표 남용 검사 — 표는 행·열 데이터에만. 기본은 도식이다 (편집 규칙)."""
    body = [sl for sec in spec["sections"] for sl in sec.get("slides", [])]
    tables = [sl for sl in body if str(sl.get("layout", "")).lower().startswith("table")]
    if body and len(tables) > max(1, len(body) // 4):
        _advise(f"본문 {len(body)}장 중 표가 {len(tables)}장이다. 행·열 데이터가 아닌 것은 "
                f"card·lead·pill·steps·compare 같은 도식 장표로 바꿀 것")

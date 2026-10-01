# -*- coding: utf-8 -*-
"""업로드 템플릿 — 내려받아 캐시하고(인스턴스 /tmp), 등록 전에는 샘플 덱 전체를 실제로 빌드해 검증한다.

패키지 tokens.py는 장표 번호·슬롯 좌표를 내장 템플릿(106장) 기준으로 갖고 있으므로, 받을 수 있는 템플릿은
**같은 구성의 새 버전**뿐이다. 장 수가 다르면 Deck()이 거절하고, 좌표가 어긋나면 샘플 빌드에서 드러난다.
"""
import tempfile
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Emu

from innogrid_ppt import tokens as T
from innogrid_ppt.template import find_template
from service import builder as B, storage as S

ROOT = Path(__file__).resolve().parent.parent
CACHE_DIR = Path(tempfile.gettempdir()) / "ppt-templates"
SAMPLE = ROOT / "sample.deck.yaml"


def fetch_template(template_url: str, template_id: str | None, work: Path) -> Path:
    """templateId가 있으면 /tmp 캐시(웜 인스턴스 재사용), 없으면 작업 폴더에 받는다."""
    if template_id and template_id.replace("-", "").isalnum():
        dest = CACHE_DIR / f"{template_id}.pptx"
        if dest.is_file() and dest.stat().st_size > 0:
            return dest
        return S.download(template_url, dest)
    return S.download(template_url, work / "template.pptx")


def _cm(v):
    return round(Emu(v).cm, 1)


def check_geometry(path: Path):
    """장 수와 슬라이드 크기가 내장 템플릿과 같아야 한다 — 패키지가 내장본 좌표(cm)로 조판하므로 세로형·다른 크기는 빌드가 되더라도 결과가 어긋난다."""
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        raise B.BuildError("template", f"pptx 파일을 열 수 없다({type(e).__name__})")
    n = len(prs.slides)
    if n != T.TEMPLATE_SLIDES:
        raise B.BuildError("template", f"템플릿이 {T.TEMPLATE_SLIDES}장이어야 하는데 {n}장이다 — 내장본과 같은 구성(장표 번호·좌표)의 새 버전만 받을 수 있다")
    base = Presentation(str(find_template(None, root=ROOT)))
    got, want = (_cm(prs.slide_width), _cm(prs.slide_height)), (_cm(base.slide_width), _cm(base.slide_height))
    if got != want:
        raise B.BuildError("template", f"슬라이드 크기가 {got[0]}×{got[1]}cm — 내장 템플릿은 {want[0]}×{want[1]}cm(가로형)다. 패키지가 내장본 좌표로 조판하므로 세로형·다른 크기 템플릿은 지원하지 않는다")


def validate_template(path: Path, work: Path) -> dict:
    """장 수·슬라이드 크기 확인 뒤 샘플 덱(모든 장표 종류)을 이 템플릿으로 빌드해 본다. 실패하면 BuildError가 올라간다."""
    check_geometry(path)
    spec = yaml.safe_load(SAMPLE.read_text(encoding="utf-8"))
    res = B.build_deck(spec, work, None, None, template=path)
    return {"ok": True, "file": path.name, "slides": res["slides"], "issues": res["issues"], "advisories": res["advisories"], "bytes": path.stat().st_size}

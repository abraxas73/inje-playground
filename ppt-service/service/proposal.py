# -*- coding: utf-8 -*-
"""제안도우미-일반현황 — 회사 일반현황 PPT를 고객사 템플릿에 맞춘다(스펙 docs/superpowers/specs/2026-10-20-proposal-general-design.md).

두 방식:
  restyle — 일반현황 장표의 도형·표·그림·차트를 벡터 그대로 옮기고, 겉모습(마스터·글꼴·브랜드 색·슬라이드 크기)만 고객사 것으로.
  reflow  — LLM이 만든 배치 JSON으로 고객사 레이아웃의 자리표시자에 제목·본문·표·그림을 새로 채운다.

패키지 innogrid_ppt는 수정하지 않는다 — transplant의 좌표 변환·그룹 도우미만 빌려 쓴다.
ponytail: 제목은 extract의 휴리스틱(상단 띠의 가장 큰 글상자)을 그대로 쓴다. 틀리면 제목이 도식에 딸려 오거나 빠진다 — 그때는 reflow로.
"""
import copy
import io
import re
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from innogrid_ppt import transplant as TR
from innogrid_ppt.clone import _clone_chart_part
from service import extract as X
from service.transplant_fix import dedupe_shape_ids, resolve_scheme_colors, theme_colors

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EMU_CM = 360000
MAX_SLIDES = 80
_TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
_BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
_SKIP_PH = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)
# 우리 브랜드 파랑(일반현황 원본) → 고객사 accent1. 의미 색(빨강·주황)은 건드리지 않는다.
BRAND_BLUES = {"006EFF", "0066FF", "0070C0", "0081FF", "0441FF", "006CDB", "2454DB"}


class ProposalError(ValueError):
    """사용자에게 보여 줄 수 있는 입력 오류(한국어)."""


# ── 파일 열기 ────────────────────────────────────────────────────

def open_pptx(path: Path) -> Presentation:
    """pptx 또는 potx(서식 파일)를 연다. potx는 본문 형식만 pptx로 바꿔 연다. 매크로 파일은 거절."""
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            names = {i.filename for i in infos}
            # 압축 폭탄 방지 — 업로드는 20MB지만 풀면 커질 수 있다
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 300 * 1024 * 1024:
                raise ProposalError("파일이 너무 큽니다(압축을 풀면 300MB 초과)")
            types = z.read("[Content_Types].xml").decode("utf-8")
    except (zipfile.BadZipFile, KeyError) as e:
        raise ProposalError("PowerPoint 파일(.pptx/.potx)이 아닙니다") from e
    if "macroEnabled" in types or any(n.endswith("vbaProject.bin") for n in names):
        raise ProposalError("매크로가 든 파일(.pptm/.potm)은 쓸 수 없습니다")
    if "presentationml.template.main+xml" in types:
        fixed = path.with_name(path.stem + "-as-pptx.pptx")
        with zipfile.ZipFile(path) as src, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                dst.writestr(item, data)
        path = fixed
    try:
        return Presentation(str(path))
    except Exception as e:  # noqa: BLE001 — 손상된 파일
        raise ProposalError("PowerPoint 파일을 열 수 없습니다(손상되었거나 형식이 다릅니다)") from e


def _clear_slides(prs):
    """템플릿에 들어 있던 예시 장표를 모두 뺀다(마스터·레이아웃은 남김)."""
    lst = prs.slides._sldIdLst
    for el in list(lst):
        prs.part.drop_rel(el.get(qn("r:id")))
        lst.remove(el)


def _cm(v):
    return round((v or 0) / EMU_CM, 2)


# ── 고객사 템플릿 분석 ──────────────────────────────────────────

def _ph_kind(ph):
    t = ph.placeholder_format.type
    if t in _TITLE_TYPES:
        return "title"
    if t == PP_PLACEHOLDER.SUBTITLE:
        return "subtitle"
    if t in _BODY_TYPES:
        return "body"
    if t == PP_PLACEHOLDER.PICTURE:
        return "picture"
    if t == PP_PLACEHOLDER.TABLE:
        return "table"
    if t in _SKIP_PH:
        return None
    return "other"


def describe_layouts(prs):
    """LLM·재구성에 쓰는 레이아웃 목록: 번호·이름·자리표시자(종류·위치 cm)."""
    out = []
    for i, layout in enumerate(prs.slide_layouts):
        phs = []
        for ph in layout.placeholders:
            kind = _ph_kind(ph)
            if kind:
                phs.append({"idx": ph.placeholder_format.idx, "kind": kind, "x": _cm(ph.left), "y": _cm(ph.top), "w": _cm(ph.width), "h": _cm(ph.height)})
        out.append({"index": i, "name": layout.name, "placeholders": phs})
    return out


def template_info(prs):
    return {"widthCm": _cm(prs.slide_width), "heightCm": _cm(prs.slide_height), "layouts": describe_layouts(prs)}


def _title_only_layout(prs):
    """제목만 있는 레이아웃(이름 'Title Only'/'제목만' 우선). 없으면 제목이 있고 자리표시자가 가장 적은 것."""
    cands = []
    for layout in prs.slide_layouts:
        kinds = [_ph_kind(p) for p in layout.placeholders]
        if "title" not in kinds:
            continue
        body = sum(1 for k in kinds if k in ("body", "picture", "table", "other", "subtitle"))
        named = bool(re.search(r"title only|제목만", layout.name or "", re.I))
        cands.append((0 if named else 1, body, layout))
    if cands:
        return min(cands, key=lambda c: (c[0], c[1]))[2]
    return prs.slide_layouts[0]


def content_area(prs, layout):
    """본문 영역 (x, y, w, h) EMU — 본문 자리표시자가 있는 레이아웃의 본문 칸, 없으면 제목 아래 여백 안쪽."""
    W, H = prs.slide_width, prs.slide_height
    for lay in [layout, *prs.slide_layouts]:
        for ph in lay.placeholders:
            if _ph_kind(ph) == "body" and ph.width and ph.height and ph.width > W * 0.5:
                return ph.left, ph.top, ph.width, ph.height
    title = next((p for p in layout.placeholders if _ph_kind(p) == "title"), None)
    top = (title.top + title.height) if title is not None and title.height else int(H * 0.18)
    return int(W * 0.05), int(top + H * 0.02), int(W * 0.90), int(H * 0.92 - top - H * 0.02)


def _set_title(slide, text):
    for ph in slide.placeholders:
        if _ph_kind(ph) == "title":
            ph.text_frame.text = text or ""
            return True
    return False


def _drop_empty_placeholders(slide):
    """채우지 않은 자리표시자는 지운다(편집 화면의 '텍스트를 입력하십시오' 상자 방지)."""
    for ph in list(slide.placeholders):
        if _ph_kind(ph) == "title":
            continue
        if ph.has_text_frame and ph.text_frame.text.strip():
            continue
        ph._element.getparent().remove(ph._element)


# ── 도형 이식(공통) ─────────────────────────────────────────────

def _materialize_placeholder(shape):
    """원본 자리표시자 도형은 위치를 레이아웃에서 물려받는다 — 좌표를 명시하고 ph 표시를 떼어 일반 도형으로."""
    el = shape._element
    nvPr = None
    for child in el.iter(qn("p:nvPr")):
        nvPr = child
        break
    if nvPr is not None:
        ph = nvPr.find(qn("p:ph"))
        if ph is not None:
            nvPr.remove(ph)
    spPr = el.find(qn("p:spPr"))
    if spPr is not None and spPr.find(qn("a:xfrm")) is None and shape.width:
        xfrm = etree.SubElement(spPr, qn("a:xfrm"))
        spPr.remove(xfrm)
        spPr.insert(0, xfrm)
        off = etree.SubElement(xfrm, qn("a:off"), x=str(shape.left), y=str(shape.top))
        ext = etree.SubElement(xfrm, qn("a:ext"), cx=str(shape.width), cy=str(shape.height))
        del off, ext


def _copy_rels(src_part, dst_part, root):
    """그림은 대상 패키지에 새 그림 파트로, 차트는 새 partname으로 복제, 외부 링크는 그대로. 나머지 관계는 버리고 참조 속성을 지운다."""
    used = set()
    for el in root.iter():
        for attr in TR._RID_ATTRS:
            v = el.get(attr)
            if v:
                used.add(v)
    mapping = {}
    for rid in used:
        rel = src_part.rels.get(rid)
        if rel is None:
            continue
        if rel.is_external:
            mapping[rid] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.IMAGE:
            _, new = dst_part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
            mapping[rid] = new
        elif rel.reltype == RT.CHART:
            mapping[rid] = dst_part.relate_to(_clone_chart_part(rel.target_part), rel.reltype)
    for el in root.iter():
        for attr in TR._RID_ATTRS:
            v = el.get(attr)
            if v is None:
                continue
            if v in mapping:
                el.set(attr, mapping[v])
            else:
                del el.attrib[attr]


def _restyle(root, scale):
    """글꼴은 고객사 테마 글꼴(+mn-lt/+mn-ea)로, 우리 브랜드 파랑은 고객사 accent1로, 줄인 만큼 글자 크기도."""
    for el in list(root.iter()):
        tag = el.tag
        if tag == qn("a:latin"):
            el.set("typeface", "+mn-lt")
        elif tag == qn("a:ea"):
            el.set("typeface", "+mn-ea")
        elif tag in (qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")):
            sz = el.get("sz")
            if sz and scale < 0.999:
                el.set("sz", str(max(600, int(int(sz) * scale))))
        elif tag == qn("a:srgbClr") and (el.get("val") or "").upper() in BRAND_BLUES:
            new = etree.Element(qn("a:schemeClr"))
            new.set("val", "accent1")
            for child in list(el):
                new.append(child)
            el.getparent().replace(el, new)


def _transplant(dst_slide, src_slide, shapes, area):
    """shapes를 한 그룹으로 묶어 area(EMU) 안에 비율 유지로 넣는다(왼쪽·위 정렬)."""
    rects = [r for r in (TR._rect(s) for s in shapes) if r]
    if not rects:
        return None
    sx, sy, sw, sh = TR.bbox(shapes)
    ax, ay, aw, ah = (v / EMU_CM for v in area)
    scale = min(aw / sw, ah / sh, 1.6) if sw > 0 and sh > 0 else 1.0
    ox, oy = ax, ay
    group = TR._new_group()
    for s in shapes:
        if s.is_placeholder:
            _materialize_placeholder(s)
        el = copy.deepcopy(s._element)
        TR._rescale(el, (sx, sy), (ox, oy), scale)
        group.append(el)
    TR._fit_group(group, ox, oy, sw * scale, sh * scale)
    dst_slide.shapes._spTree.append(group)
    resolve_scheme_colors(group, theme_colors(src_slide))
    _copy_rels(src_slide.part, dst_slide.part, group)
    _restyle(group, scale)
    dedupe_shape_ids(dst_slide, group)
    return group


# ── restyle ─────────────────────────────────────────────────────

def _body_shapes(src_slide, title_bottom_cm, title_text):
    """제목 띠 아래 최상위 도형(제목 글상자·바닥글 자리표시자 제외)."""
    out = []
    for s in src_slide.shapes:
        if s.is_placeholder and s.placeholder_format.type in (*_SKIP_PH, *_TITLE_TYPES):
            continue
        r = TR._rect(s)
        if r is None:
            continue
        if title_bottom_cm is not None and r[1] + r[3] <= title_bottom_cm + 0.01:
            continue
        if title_text and getattr(s, "has_text_frame", False) and s.has_text_frame and s.text_frame.text.strip().startswith(title_text):
            continue
        out.append(s)
    return out


def restyle(profile_path: Path, template_path: Path, out_path: Path) -> dict:
    src = open_pptx(profile_path)
    dst = open_pptx(template_path)
    if len(src.slides) > MAX_SLIDES:
        raise ProposalError(f"일반현황 PPT는 {MAX_SLIDES}장까지입니다")
    outline = X.extract_slides(str(profile_path))
    _clear_slides(dst)
    layout = _title_only_layout(dst)
    area = content_area(dst, layout)
    moved = 0
    for info, src_slide in zip(outline, src.slides):
        slide = dst.slides.add_slide(layout)
        title = next((p.text_frame.text for p in src_slide.placeholders if _ph_kind(p) == "title" and p.text_frame.text.strip()), None) or info["title"] or ""
        _set_title(slide, title)
        shapes = _body_shapes(src_slide, info["titleBottomCm"], title.strip())
        if shapes and _transplant(slide, src_slide, shapes, area) is not None:
            moved += len(shapes)
        _drop_empty_placeholders(slide)
    dst.save(str(out_path))
    return {"slides": len(dst.slides), "shapes": moved}


# ── reflow ──────────────────────────────────────────────────────

def profile_outline(prs):
    """LLM 재료: 장표별 제목·글 상자·표(셀 글)·그림 번호. 그림 id는 's{장}-{n}'(reflow가 원본에서 다시 찾는다)."""
    out = []
    for no, slide in enumerate(prs.slides, 1):
        texts, tables, images = [], [], []
        title = None
        for s in X._walk(slide.shapes):
            if s.is_placeholder and s.placeholder_format.type in _TITLE_TYPES and s.text_frame.text.strip():
                title = s.text_frame.text.strip()
                continue
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images.append({"id": f"s{no}-{len(images) + 1}", "w": _cm(s.width), "h": _cm(s.height)})
            elif getattr(s, "has_table", False) and s.has_table:
                tables.append([[c.text.strip()[:200] for c in row.cells] for row in s.table.rows][:40])
            elif getattr(s, "has_text_frame", False) and s.has_text_frame and s.text_frame.text.strip():
                texts.append(s.text_frame.text.strip()[:1500])
        out.append({"no": no, "title": title, "texts": texts, "tables": tables, "images": images})
    return out


def _source_images(prs):
    found = {}
    for no, slide in enumerate(prs.slides, 1):
        n = 0
        for s in X._walk(slide.shapes):
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n += 1
                found[f"s{no}-{n}"] = s.image.blob
    return found


def _fill_body(ph, items):
    tf = ph.text_frame
    tf.clear()
    for i, item in enumerate(items):
        text, level = (item.get("text", ""), int(item.get("level", 0))) if isinstance(item, dict) else (str(item), 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = max(0, min(4, level))
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for old in bodyPr.findall(qn("a:normAutofit")) + bodyPr.findall(qn("a:spAutoFit")):
        bodyPr.remove(old)
    etree.SubElement(bodyPr, qn("a:normAutofit"))  # 넘치면 PowerPoint가 글자를 줄인다


def _add_table(slide, rect, table):
    header, rows = table.get("header") or [], table.get("rows") or []
    data = ([header] if header else []) + rows
    if not data:
        return
    ncol = max(len(r) for r in data)
    x, y, w, h = rect
    row_h = int(min(h / len(data), Emu(Pt(28))))
    shape = slide.shapes.add_table(len(data), ncol, x, y, w, row_h * len(data))
    size = Pt(14 if len(data) <= 8 else 12 if len(data) <= 14 else 10)
    for r, row in enumerate(data):
        for c in range(ncol):
            cell = shape.table.cell(r, c)
            cell.text = str(row[c]) if c < len(row) else ""
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = size
    shape.table.first_row = bool(header)


def _rect_of(ph):
    return ph.left, ph.top, ph.width, ph.height


def reflow(profile_path: Path, template_path: Path, spec: dict, out_path: Path) -> dict:
    src = open_pptx(profile_path)
    dst = open_pptx(template_path)
    slides = spec.get("slides") if isinstance(spec, dict) else None
    if not isinstance(slides, list) or not slides or len(slides) > MAX_SLIDES:
        raise ProposalError("배치 결과가 비었거나 너무 깁니다")
    images = _source_images(src)
    _clear_slides(dst)
    layouts = list(dst.slide_layouts)
    for s in slides:
        idx = s.get("layout", 0)
        layout = layouts[idx] if isinstance(idx, int) and 0 <= idx < len(layouts) else _title_only_layout(dst)
        slide = dst.slides.add_slide(layout)
        _set_title(slide, str(s.get("title") or ""))
        bodies = [p for p in slide.placeholders if _ph_kind(p) == "body"]
        pics = [p for p in slide.placeholders if _ph_kind(p) == "picture"]
        area = content_area(dst, layout)
        table = s.get("table") if isinstance(s.get("table"), dict) else None
        body = s.get("body") if isinstance(s.get("body"), list) else []
        img_ids = [i for i in (s.get("images") or []) if i in images][:4]
        # 표가 있으면 본문 칸(없으면 본문 영역)에 그리고, 글은 두 번째 본문 칸이 있을 때만 그 칸에
        if table:
            target = bodies.pop(0) if bodies else None
            rect = _rect_of(target) if target is not None else area
            if target is not None:
                target._element.getparent().remove(target._element)
            if body and bodies:
                _fill_body(bodies.pop(0), body)
                body = []
            _add_table(slide, rect, table)
        if body:
            if bodies:
                _fill_body(bodies.pop(0), body)
            else:
                box = slide.shapes.add_textbox(*area)
                _fill_body(box, body)
        for i, img in enumerate(img_ids):
            blob = io.BytesIO(images[img])
            if pics:
                pics.pop(0).insert_picture(blob)
            else:
                x, y, w, h = area
                cw = w // max(1, len(img_ids))
                pic = slide.shapes.add_picture(blob, x + cw * i, y)
                ratio = min(cw / pic.width, h / pic.height, 1.0)
                pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
        _drop_empty_placeholders(slide)
    dst.save(str(out_path))
    return {"slides": len(dst.slides)}


def cleanup(path: Path):
    shutil.rmtree(path, ignore_errors=True)
